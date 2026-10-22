"""Data layer for CineScope.

Everything that touches the Excel file lives here. data/movies.xlsx is one
workbook: the "movies" sheet is the web's data, the other sheets hold the
group's Analysis ToolPak results.
- first-run conversion (from data/movies.csv, or a raw "movies" sheet)
- data cleaning (rating merge, splitting the "released" column)
- reading the workbook (with a small cache keyed on file modification time)
- safe writing that replaces only the "movies" sheet (temp file, then replace)
- add / update / delete a movie, and server-side form validation
- reading the Analysis ToolPak results (never written by the web)

Statistics and models are NOT here; they belong in analysis.py.
"""

import os
from datetime import date, datetime

import openpyxl
import pandas as pd

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
CSV_PATH = os.path.join(DATA_DIR, "movies.csv")
XLSX_PATH = os.path.join(DATA_DIR, "movies.xlsx")
TEMP_PATH = os.path.join(DATA_DIR, "movies.tmp.xlsx")
SHEET_NAME = "movies"

# Column order stored in the Excel file.
COLUMNS = [
    "id", "name", "rating", "genre", "year", "release_date", "release_country",
    "score", "votes", "director", "writer", "star", "country",
    "budget", "gross", "company", "runtime",
]
TEXT_COLUMNS = [
    "name", "rating", "genre", "release_date", "release_country",
    "director", "writer", "star", "country", "company",
]
NUMBER_COLUMNS = ["year", "score", "votes", "budget", "gross", "runtime"]

# Ratings that mean the same thing are merged into one label.
RATING_ALIASES = {"Unrated": "Not Rated"}


class ExcelLockedError(Exception):
    """Raised when the Excel file cannot be written (usually open in Excel)."""

    def __init__(self):
        super().__init__(
            "บันทึกข้อมูลไม่สำเร็จ เนื่องจากไฟล์ data/movies.xlsx ถูกเปิดค้างอยู่ "
            "(เช่น เปิดในโปรแกรม Excel) กรุณาปิดไฟล์แล้วลองใหม่อีกครั้ง"
        )


class MovieNotFoundError(Exception):
    """Raised when a movie id does not exist."""

    def __init__(self, movie_id):
        super().__init__(f"ไม่พบภาพยนตร์รหัส {movie_id}")


# ---------------------------------------------------------------------------
# Cleaning
# ---------------------------------------------------------------------------

def normalize_rating(value):
    """Merge duplicate rating labels, e.g. "Unrated" -> "Not Rated"."""
    if pd.isna(value):
        return value
    value = str(value).strip()
    return RATING_ALIASES.get(value, value)


def split_released(released):
    """Split "June 13, 1980 (United States)" into ("1980-06-13", "United States").

    Full dates become ISO format (YYYY-MM-DD). Partial dates such as
    "November 1980" are kept as written so no information is invented.
    """
    released = released.fillna("").astype(str)
    parts = released.str.extract(r"^(.*?)\s*\((.*)\)\s*$")
    date_text = parts[0].fillna(released).str.strip()
    country = parts[1].str.strip()

    parsed = pd.to_datetime(date_text, format="%B %d, %Y", errors="coerce")
    release_date = parsed.dt.strftime("%Y-%m-%d").where(parsed.notna(), date_text)
    release_date = release_date.replace("", pd.NA)
    return release_date, country


def clean_raw_csv(raw):
    """Turn the original CSV table into the table stored in Excel."""
    df = raw.copy()
    df["rating"] = df["rating"].map(normalize_rating)
    df["release_date"], df["release_country"] = split_released(df["released"])
    df = df.drop(columns=["released"])
    df.insert(0, "id", range(1, len(df) + 1))
    return df[COLUMNS]


def add_financials(df):
    """Add profit and roi columns.

    Only rows with both budget and gross (and budget > 0) get values;
    the rest stay empty so financial analysis can simply drop them.
    """
    has_money = df["budget"].notna() & df["gross"].notna() & (df["budget"] > 0)
    df["profit"] = (df["gross"] - df["budget"]).where(has_money)
    df["roi"] = (df["gross"] / df["budget"]).where(has_money)
    return df


def prepare_loaded(df):
    """Fix types and derived columns after reading the workbook."""
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = pd.NA
    df = df[COLUMNS].copy()

    df["id"] = pd.to_numeric(df["id"], errors="coerce").astype("Int64")
    for col in NUMBER_COLUMNS:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    for col in TEXT_COLUMNS:
        df[col] = df[col].astype("string").str.strip().replace("", pd.NA)

    df["rating"] = df["rating"].map(normalize_rating).astype("string")
    df = df[df["id"].notna()]
    return add_financials(df)


# ---------------------------------------------------------------------------
# Reading / writing Excel
# ---------------------------------------------------------------------------

# data/movies.xlsx is ONE workbook shared with the group's Excel work:
#   - sheet "movies": the movie table the web reads and writes
#   - every other sheet (Analysis ToolPak results, ...) is kept untouched
# Note: openpyxl cannot keep Excel charts/images, so do not put charts in
# this workbook; the ToolPak result tables (plain values) are kept safely.

_cache = {"mtime": None, "df": None}
_checked = {"mtime": None}


def _movies_sheet_header():
    """Header row of the "movies" sheet, or None when the sheet does not exist."""
    workbook = openpyxl.load_workbook(XLSX_PATH, read_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            return None
        first = next(workbook[SHEET_NAME].iter_rows(max_row=1, values_only=True), ())
        return [str(v).strip() for v in first if v is not None]
    finally:
        workbook.close()


def _read_csv():
    if not os.path.exists(CSV_PATH):
        raise FileNotFoundError(
            f"ไม่พบไฟล์ข้อมูล {CSV_PATH} กรุณาวาง movies.csv ไว้ในโฟลเดอร์ data/"
        )
    return pd.read_csv(CSV_PATH)


def ensure_excel():
    """Make sure data/movies.xlsx has a "movies" sheet in the web's format.

    - no workbook yet          -> create it from data/movies.csv
    - workbook without "movies" -> add the sheet from data/movies.csv
    - "movies" is raw data      -> clean it once (add id, split "released", merge ratings)
    Other sheets are always kept. Returns True when something was converted.
    """
    if not os.path.exists(XLSX_PATH):
        write_excel(clean_raw_csv(_read_csv()))
        return True
    mtime = os.path.getmtime(XLSX_PATH)
    if _checked["mtime"] == mtime:  # already checked this version of the file
        return False

    header = _movies_sheet_header()
    converted = False
    if header is None:
        write_excel(clean_raw_csv(_read_csv()))
        converted = True
    elif "id" not in header and "released" in header:
        raw = pd.read_excel(XLSX_PATH, sheet_name=SHEET_NAME, engine="openpyxl")
        write_excel(clean_raw_csv(raw))
        converted = True
    _checked["mtime"] = os.path.getmtime(XLSX_PATH)
    return converted


def read_excel():
    """Read the "movies" sheet from disk (no cache, no derived columns)."""
    return pd.read_excel(XLSX_PATH, sheet_name=SHEET_NAME, engine="openpyxl")


def _cell(value):
    """Value for an Excel cell: None for missing, plain Python types for numpy."""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    return value.item() if hasattr(value, "item") else value


def write_excel(df):
    """Write the movie table safely, keeping every other sheet in the workbook.

    Only the "movies" sheet is replaced (at the same position). The workbook is
    saved to a temporary file first and then moved over the real file, so a
    crash mid-write never leaves a half-written workbook.
    """
    try:
        if os.path.exists(XLSX_PATH):
            workbook = openpyxl.load_workbook(XLSX_PATH)
            position = 0
            if SHEET_NAME in workbook.sheetnames:
                position = workbook.sheetnames.index(SHEET_NAME)
                del workbook[SHEET_NAME]
            sheet = workbook.create_sheet(SHEET_NAME, position)
        else:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = SHEET_NAME
        sheet.append(COLUMNS)
        for row in df[COLUMNS].itertuples(index=False):
            sheet.append([_cell(v) for v in row])
        workbook.save(TEMP_PATH)
        os.replace(TEMP_PATH, XLSX_PATH)
    except PermissionError as exc:
        _remove_temp_file()
        raise ExcelLockedError() from exc
    except Exception:
        _remove_temp_file()
        raise
    _cache["mtime"] = None  # force a fresh read next time


def _remove_temp_file():
    try:
        if os.path.exists(TEMP_PATH):
            os.remove(TEMP_PATH)
    except OSError:
        pass


def data_version():
    """A value that changes whenever the Excel file changes (for model caching)."""
    ensure_excel()
    return os.path.getmtime(XLSX_PATH)


def load_movies():
    """Return the cleaned movie table, including profit and roi.

    The workbook is re-read only when its modification time changes, so pages
    always reflect the current file without re-parsing it on every request.
    A copy is returned so callers cannot modify the cache.
    """
    mtime = data_version()
    if _cache["df"] is None or _cache["mtime"] != mtime:
        _cache["df"] = prepare_loaded(read_excel())
        _cache["mtime"] = mtime
    return _cache["df"].copy()


# ---------------------------------------------------------------------------
# CRUD
# ---------------------------------------------------------------------------

def get_movie(movie_id):
    """Return one movie as a dict (missing values become None)."""
    df = load_movies()
    row = df[df["id"] == int(movie_id)]
    if row.empty:
        raise MovieNotFoundError(movie_id)
    record = row.iloc[0].to_dict()
    return {key: (None if pd.isna(value) else value) for key, value in record.items()}


def add_movie(record):
    """Append a validated record and return its new id."""
    df = load_movies()
    new_id = int(df["id"].max()) + 1 if len(df) else 1
    new_row = {col: record.get(col) for col in COLUMNS}
    new_row["id"] = new_id
    df = pd.concat([df[COLUMNS], pd.DataFrame([new_row])], ignore_index=True)
    write_excel(df)
    return new_id


def update_movie(movie_id, record):
    """Overwrite the fields of an existing movie with a validated record."""
    df = load_movies()[COLUMNS]
    mask = df["id"] == int(movie_id)
    if not mask.any():
        raise MovieNotFoundError(movie_id)
    for col in COLUMNS:
        if col != "id":
            df.loc[mask, col] = record.get(col)
    write_excel(df)


def delete_movie(movie_id):
    """Remove a movie by id."""
    df = load_movies()[COLUMNS]
    mask = df["id"] == int(movie_id)
    if not mask.any():
        raise MovieNotFoundError(movie_id)
    write_excel(df[~mask])


def get_choices(df=None):
    """Distinct genre, rating and country values for form dropdowns."""
    if df is None:
        df = load_movies()

    def distinct(col):
        return sorted(df[col].dropna().astype(str).unique().tolist())

    return {
        "genre": distinct("genre"),
        "rating": distinct("rating"),
        "country": distinct("country"),
        "release_country": distinct("release_country"),
    }


# ---------------------------------------------------------------------------
# Search / filter / sort / paginate (used by the movie list page)
# ---------------------------------------------------------------------------

SORTABLE_COLUMNS = ["year", "score", "budget", "gross"]


def search_movies(df, q="", genre="", rating="", year_min=None, year_max=None,
                  sort="year", order="desc", page=1, per_page=20):
    """Filter, sort and paginate the movie table.

    Returns a dict with the rows of the requested page (as dicts) plus
    the total number of matches and the number of pages.
    """
    result = df
    if q:
        text = q.strip().lower()
        match = (
            result["name"].str.lower().str.contains(text, regex=False, na=False)
            | result["director"].str.lower().str.contains(text, regex=False, na=False)
            | result["star"].str.lower().str.contains(text, regex=False, na=False)
        )
        result = result[match]
    if genre:
        result = result[result["genre"] == genre]
    if rating:
        result = result[result["rating"] == rating]
    if year_min is not None:
        result = result[result["year"] >= year_min]
    if year_max is not None:
        result = result[result["year"] <= year_max]

    if sort not in SORTABLE_COLUMNS:
        sort = "year"
    ascending = order == "asc"
    # Missing values always go to the end; "name" breaks ties for a stable order.
    result = result.sort_values([sort, "name"], ascending=[ascending, True],
                                na_position="last", kind="mergesort")

    total = len(result)
    pages = max(1, -(-total // per_page))  # ceiling division
    page = min(max(1, page), pages)
    start = (page - 1) * per_page
    rows = result.iloc[start:start + per_page]

    items = [
        {key: (None if pd.isna(value) else value) for key, value in row.items()}
        for row in rows.to_dict(orient="records")
    ]
    return {"items": items, "total": total, "page": page, "pages": pages}


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def _parse_number(form, field, errors, label, minimum=None, maximum=None,
                  integer=False, required=False):
    text = str(form.get(field, "") or "").strip().replace(",", "")
    if text == "":
        if required:
            errors[field] = f"กรุณากรอก{label}"
        return None
    try:
        value = float(text)
    except ValueError:
        errors[field] = f"{label}ต้องเป็นตัวเลข"
        return None
    if integer and not value.is_integer():
        errors[field] = f"{label}ต้องเป็นจำนวนเต็ม"
        return None
    if minimum is not None and value < minimum:
        errors[field] = _range_message(label, minimum, maximum)
        return None
    if maximum is not None and value > maximum:
        errors[field] = _range_message(label, minimum, maximum)
        return None
    return int(value) if integer else value


def _range_message(label, minimum, maximum):
    if maximum is None:
        return f"{label}ต้องไม่ติดลบ" if minimum == 0 else f"{label}ต้องไม่น้อยกว่า {minimum:g}"
    return f"{label}ต้องอยู่ระหว่าง {minimum:g} ถึง {maximum:g}"


def is_valid_release_date(text):
    """Accept a full ISO date, or the partial forms found in the dataset."""
    for fmt in ("%Y-%m-%d", "%B %Y", "%Y"):
        try:
            datetime.strptime(text, fmt)
            return True
        except ValueError:
            continue
    return False


def validate_movie(form):
    """Validate submitted form data.

    Returns (record, errors). `record` holds cleaned values ready to save;
    `errors` maps field name -> Thai error message (empty when valid).
    """
    errors = {}
    record = {}

    for field in TEXT_COLUMNS:
        text = str(form.get(field, "") or "").strip()
        record[field] = text or None

    if not record["name"]:
        errors["name"] = "กรุณากรอกชื่อเรื่อง"
    if not record["genre"]:
        errors["genre"] = "กรุณาเลือกประเภทหนัง"
    record["rating"] = normalize_rating(record["rating"]) if record["rating"] else None

    if record["release_date"] and not is_valid_release_date(record["release_date"]):
        errors["release_date"] = (
            "รูปแบบวันฉายไม่ถูกต้อง ใช้ ปี-เดือน-วัน เช่น 1980-06-13 "
            "(หรือ November 1980 / 1980 ถ้าไม่ทราบวันที่)"
        )

    current_year = date.today().year
    record["year"] = _parse_number(form, "year", errors, "ปีที่ฉาย",
                                   1900, current_year, integer=True, required=True)
    record["score"] = _parse_number(form, "score", errors, "คะแนน", 0, 10)
    record["votes"] = _parse_number(form, "votes", errors, "จำนวนโหวต", 0, integer=True)
    record["budget"] = _parse_number(form, "budget", errors, "ทุนสร้าง", 0)
    record["gross"] = _parse_number(form, "gross", errors, "รายได้", 0)
    record["runtime"] = _parse_number(form, "runtime", errors, "ความยาว", 1, 600)

    return record, errors


# ---------------------------------------------------------------------------
# Analysis ToolPak results (other sheets of data/movies.xlsx) - READ ONLY
# ---------------------------------------------------------------------------
# The group runs Excel's Analysis ToolPak inside data/movies.xlsx. The web only
# reads these sheets; write_excel() replaces the "movies" sheet and keeps them.
# Values are found by their row labels ("R Square", "Mean", ...), not by fixed
# cell addresses, so the layout can shift a little without breaking anything.

TOOLPAK_SHEETS = {
    "descriptive_prefix": "Descriptive",   # every sheet starting with this
    "correlation": "Correlation",
    "regression": "Regression",
    "regression_data": "Regression_Data",  # the X/Y columns given to Regression
}

_toolpak_cache = {"mtime": None, "data": None}


def _sheet_rows(workbook, name):
    """All rows of a sheet as lists of values (empty list if the sheet is missing)."""
    if name not in workbook.sheetnames:
        return []
    return [list(row) for row in workbook[name].iter_rows(values_only=True)]


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _rows_to_frame(rows):
    """First row = header; drops fully empty rows."""
    if not rows:
        return pd.DataFrame()
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    body = [r for r in rows[1:] if any(v is not None for v in r)]
    return pd.DataFrame(body, columns=header)


def parse_descriptive(rows):
    """ToolPak Descriptive Statistics: variable names on the first row, then
    label/value pairs (A/B, C/D, ...) below. Returns {variable: {label: value}}."""
    result = {}
    if not rows:
        return result
    for col in range(0, len(rows[0]), 2):
        name = rows[0][col]
        if name is None:
            continue
        stats = {}
        for row in rows[1:]:
            label = row[col] if col < len(row) else None
            value = row[col + 1] if col + 1 < len(row) else None
            if label is not None and _is_number(value):
                stats[str(label).strip()] = float(value)
        result[str(name).strip()] = stats
    return result


def parse_correlation(rows):
    """ToolPak Correlation matrix (lower triangle). Returns
    {"labels": [...], "pairs": {(a, b): r}} with every pair stored once."""
    if not rows:
        return {"labels": [], "pairs": {}}
    labels = [str(v).strip() for v in rows[0][1:] if v is not None]
    pairs = {}
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        a = str(row[0]).strip()
        for j, b in enumerate(labels):
            value = row[j + 1] if j + 1 < len(row) else None
            if _is_number(value) and a != b and (b, a) not in pairs:
                pairs[(a, b)] = float(value)
    return {"labels": labels, "pairs": pairs}


def parse_regression(rows):
    """ToolPak Regression SUMMARY OUTPUT: statistics, ANOVA and coefficients."""
    stats, anova, coefficients, headers = {}, {}, [], []
    section = "stats"
    for row in rows:
        first = str(row[0]).strip() if row and row[0] is not None else ""
        second = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if first == "ANOVA":
            section = "anova"
            continue
        if second == "Coefficients":
            section = "coefficients"
            headers = [str(v).strip() if v is not None else "" for v in row[1:]]
            continue
        if not first:
            continue
        if section == "stats" and len(row) > 1 and _is_number(row[1]):
            stats[first] = float(row[1])
        elif section == "anova":
            anova[first] = {h: float(v) for h, v in zip(["df", "SS", "MS", "F", "Significance F"], row[1:6])
                            if _is_number(v)}
        elif section == "coefficients":
            values = {h: float(v) for h, v in zip(headers, row[1:]) if h and _is_number(v)}
            if values:
                coefficients.append({"name": first, **values})
    return {"stats": stats, "anova": anova, "coefficients": coefficients, "headers": headers}


def read_toolpak_results():
    """Read the ToolPak sheets of data/movies.xlsx. Returns None when there are
    no ToolPak results yet, a dict with "error" when they cannot be read,
    otherwise the parsed results. Re-read only when the file changes."""
    ensure_excel()
    mtime = os.path.getmtime(XLSX_PATH)
    if _toolpak_cache["mtime"] == mtime:
        return _toolpak_cache["data"]

    try:
        workbook = openpyxl.load_workbook(XLSX_PATH, data_only=True, read_only=True)
        try:
            descriptive = {}
            for name in workbook.sheetnames:
                if name.startswith(TOOLPAK_SHEETS["descriptive_prefix"]):
                    descriptive.update(parse_descriptive(_sheet_rows(workbook, name)))
            data = {
                "sheets": workbook.sheetnames,
                "descriptive": descriptive,
                "correlation": parse_correlation(_sheet_rows(workbook, TOOLPAK_SHEETS["correlation"])),
                "regression": parse_regression(_sheet_rows(workbook, TOOLPAK_SHEETS["regression"])),
                "regression_data": _rows_to_frame(_sheet_rows(workbook, TOOLPAK_SHEETS["regression_data"])),
            }
        finally:
            workbook.close()
        if not (data["descriptive"] or data["correlation"]["pairs"] or data["regression"]["stats"]):
            data = None  # the workbook has no ToolPak sheets yet
    except Exception as exc:  # a broken or locked file must not break the page
        data = {"error": f"อ่านผล Analysis ToolPak ใน movies.xlsx ไม่ได้ ({exc.__class__.__name__}: {exc})"}

    _toolpak_cache["mtime"] = mtime
    _toolpak_cache["data"] = data
    return data
